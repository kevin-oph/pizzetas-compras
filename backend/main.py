from fastapi import FastAPI, Depends, HTTPException, status, Query, UploadFile, File, Form
from sqlalchemy.orm import Session
from sqlalchemy import func
from fastapi.middleware.cors import CORSMiddleware
from typing import List, Dict, Optional, Any
from pydantic import BaseModel
from datetime import datetime, timedelta, timezone
import io
import csv
import openpyxl

# Importaciones locales
import models
from database import engine, get_db

# Crea las tablas en la BD si no existen
models.Base.metadata.create_all(bind=engine)

app = FastAPI(title="Pizzetas Artesanales API - Sistema Integral 360°")

# Configuración de CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

CATEGORIES_LIST = [
    "Carnes",
    "Quesos y Lácteos",
    "Vinos y Licores",
    "Abarrotes y Harinas",
    "Desechables y Empaque",
    "Limpieza"
]

def get_stock_status(current: float, ideal: float):
    if current <= 0:
        return "Critical"
    elif current < ideal:
        return "Low"
    return "Optimal"

# --- ESQUEMAS PYDANTIC ---
class StockConsume(BaseModel):
    product_id: int
    consumed_amount: float
    reason_type: Optional[str] = "MERMA" # MERMA, LIMPIEZA, COMIDA_PERSONAL, OTRO
    shift_notes: Optional[str] = "Salida extraordinaria"
    user_id: Optional[int] = None

class StockUpdate(BaseModel):
    product_id: int
    new_stock: float

class PurchaseItemIn(BaseModel):
    product_id: int
    quantity: float
    unit_price_paid: float

class PurchaseCreate(BaseModel):
    provider_id: int
    ticket_number: Optional[str] = None
    notes: Optional[str] = None
    user_id: Optional[int] = None
    items: List[PurchaseItemIn]

class ProviderCreate(BaseModel):
    name: str

class ProductCreate(BaseModel):
    name: str
    category: str = "Abarrotes y Harinas"
    provider_id: int
    ideal_stock: float
    current_stock: float = 0.0
    unit_measure: str
    unit_price: float

class ProductUpdate(BaseModel):
    name: str
    category: str = "Abarrotes y Harinas"
    ideal_stock: float
    unit_measure: str
    unit_price: float
    provider_id: Optional[int] = None

class ParseTextRequest(BaseModel):
    text: str

class RecipeItemIn(BaseModel):
    product_id: int
    quantity_needed: float

class MenuItemCreate(BaseModel):
    pos_name: str
    category: str = "COMIDA"
    sale_price: float = 0.0
    recipes: List[RecipeItemIn] = []

class ConfirmCutItemIn(BaseModel):
    name: str
    quantity: float
    total_amount: float
    unit_price: float = 0.0
    cogs_cost: float = 0.0
    class Config:
        extra = "ignore"

class ConfirmCutIngredientIn(BaseModel):
    product_id: int
    quantity_to_deduct: float
    cost: Optional[float] = 0.0
    total_cost: Optional[float] = 0.0
    class Config:
        extra = "ignore"

class ConfirmCutRequest(BaseModel):
    cut_date_str: str
    total_sales: float
    cash_sales: float = 0.0
    card_sales: float = 0.0
    transfer_sales: float = 0.0
    tips_cash: float = 0.0
    tips_card: float = 0.0
    tips_total: float = 0.0
    cash_balance: float = 0.0
    total_items_sold: int = 0
    total_cogs: float = 0.0
    gross_profit: float = 0.0
    food_cost_percentage: float = 0.0
    sold_items: List[ConfirmCutItemIn] = []
    ingredients_to_deduct: List[ConfirmCutIngredientIn] = []
    class Config:
        extra = "ignore"


# --- UTILIDADES DE PARSEO DE CORTE POS ---

def parse_corte_rows(rows: List[List[str]]):
    header = {
        "date_str": "Fecha no detectada",
        "efectivo": 0.0,
        "tarjeta": 0.0,
        "transferencias": 0.0,
        "ventas_totales": 0.0,
        "propinas_efectivo": 0.0,
        "propinas_tarjeta": 0.0,
        "propinas_totales": 0.0,
        "saldo_caja": 0.0,
    }
    
    products = []
    section = None
    
    for row in rows:
        cleaned = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
        if not cleaned:
            continue
            
        row_str = " ".join(cleaned)
        
        # Detectar fecha
        if "Fecha:" in row_str:
            for idx, c in enumerate(cleaned):
                if c == "Fecha:" and idx + 1 < len(cleaned):
                    header["date_str"] = cleaned[idx + 1]
                    break
                elif any(m in c.lower() for m in ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]):
                    header["date_str"] = c
                    break
        
        # Detectar pares clave-valor de caja en toda la fila (puede haber múltiples columnas)
        for i in range(len(cleaned) - 1):
            key = cleaned[i].lower()
            val_str = cleaned[i + 1].replace("$", "").replace(",", "").strip()
            try:
                val = float(val_str)
                if key == "efectivo" and header["efectivo"] == 0:
                    header["efectivo"] = val
                elif key == "tarjeta" and header["tarjeta"] == 0:
                    header["tarjeta"] = val
                elif key == "transferencias" and header["transferencias"] == 0:
                    header["transferencias"] = val
                elif key == "ventastotales":
                    header["ventas_totales"] = val
                elif key == "propinas efectivo":
                    header["propinas_efectivo"] = val
                elif key == "propinas tarjeta":
                    header["propinas_tarjeta"] = val
                elif key == "propinas totales":
                    header["propinas_totales"] = val
                elif key == "saldo en caja":
                    header["saldo_caja"] = val
            except ValueError:
                pass

        # Detectar sección de productos
        if len(cleaned) >= 2 and cleaned[0].lower() == "producto" and "cantidad" in cleaned[1].lower():
            section = "PRODUCTS"
            continue
        elif cleaned[0].lower() in ["folio", "categoría", "categoria", "tickets de ventas"]:
            section = "OTHER"
            continue
        elif section == "PRODUCTS":
            if len(cleaned) >= 3:
                name = cleaned[0].strip()
                try:
                    qty = float(cleaned[1].replace(",", ""))
                    total = float(cleaned[2].replace(",", ""))
                    unit_p = round(total / qty, 2) if qty > 0 else 0.0
                    products.append({
                        "name": name,
                        "quantity": qty,
                        "total_amount": total,
                        "unit_price": unit_p
                    })
                except ValueError:
                    pass

    # Si ventas_totales fue 0 pero hay desglose de pagos, sumar
    if header["ventas_totales"] == 0:
        header["ventas_totales"] = header["efectivo"] + header["tarjeta"] + header["transferencias"]
    if header["ventas_totales"] == 0 and products:
        header["ventas_totales"] = sum(p["total_amount"] for p in products)

    return header, products


# --- ENDPOINTS BÁSICOS Y CATEGORÍAS ---

@app.get("/")
def read_root():
    return {"message": "API Pizzetas Artesanales 360° Activa con Integración POS", "version": "3.0"}

@app.get("/api/categories")
def get_categories(db: Session = Depends(get_db)):
    """Lista las categorías con conteo de productos."""
    counts = db.query(models.Product.category, func.count(models.Product.id)).group_by(models.Product.category).all()
    count_map = {cat: cnt for cat, cnt in counts}
    
    result = []
    for cat in CATEGORIES_LIST:
        result.append({
            "name": cat,
            "product_count": count_map.get(cat, 0)
        })
    return result


# --- CATÁLOGO DE INVENTARIO ---

@app.get("/api/inventory-catalog")
def get_inventory_catalog(
    category: Optional[str] = None, 
    provider_id: Optional[int] = None, 
    db: Session = Depends(get_db)
):
    query = db.query(models.Product)
    if category and category != "ALL":
        query = query.filter(models.Product.category == category)
    if provider_id:
        query = query.filter(models.Product.provider_id == provider_id)
        
    products = query.order_by(models.Product.name.asc()).all()
    catalog = {}
    
    for product in products:
        provider_name = product.provider.name if product.provider else "Sin Proveedor"
        if provider_name not in catalog:
            catalog[provider_name] = []
            
        catalog[provider_name].append({
            "product_id": product.id,
            "name": product.name,
            "category": product.category or "Abarrotes y Harinas",
            "current_stock": product.current_stock or 0.0,
            "ideal_stock": product.ideal_stock or 0.0,
            "unit": product.unit_measure,
            "unit_price": product.unit_price or 0.0,
            "last_purchased_price": product.last_purchased_price or product.unit_price or 0.0,
            "last_purchased_date": product.last_purchased_date.strftime("%Y-%m-%d %H:%M") if product.last_purchased_date else None,
            "status": get_stock_status(product.current_stock or 0.0, product.ideal_stock or 0.0)
        })
        
    return catalog


# --- LISTA Y CHECKLIST DE COMPRAS ---

@app.get("/api/shopping-list")
def generate_shopping_list(category: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(models.Product)
    if category and category != "ALL":
        query = query.filter(models.Product.category == category)
        
    products = query.all()
    
    shopping_data = {}
    for product in products:
        stock = product.current_stock or 0.0
        ideal = product.ideal_stock or 0.0
        
        if stock < ideal:
            provider_name = product.provider.name if product.provider else "Sin Proveedor"
            provider_id = product.provider.id if product.provider else 0
            
            if provider_name not in shopping_data:
                shopping_data[provider_name] = {
                    "provider_id": provider_id,
                    "items": [],
                    "estimated_provider_cost": 0.0
                }
            
            buy_amount = round(ideal - stock, 2)
            unit_price = product.unit_price or 0.0
            total_estimated_price = round(buy_amount * unit_price, 2)
            
            shopping_data[provider_name]["items"].append({
                "product_id": product.id,
                "name": product.name,
                "category": product.category or "Abarrotes y Harinas",
                "current_stock": stock,
                "ideal_stock": ideal,
                "buy_amount": buy_amount,
                "unit": product.unit_measure,
                "unit_price": unit_price,
                "last_purchased_price": product.last_purchased_price or unit_price,
                "total_estimated_price": total_estimated_price
            })
            
            shopping_data[provider_name]["estimated_provider_cost"] += total_estimated_price
        
    final_shopping_data = {
        prov: {
            **data,
            "estimated_provider_cost": round(data["estimated_provider_cost"], 2)
        } 
        for prov, data in shopping_data.items() if len(data["items"]) > 0
    }
    return final_shopping_data


# --- REGISTRO DE COMPRAS REALES Y TICKETS ---

@app.post("/api/register-purchase")
def register_purchase(data: PurchaseCreate, db: Session = Depends(get_db)):
    provider = db.query(models.Provider).filter(models.Provider.id == data.provider_id).first()
    if not provider:
        raise HTTPException(status_code=404, detail="Proveedor no encontrado")

    if not data.items:
        raise HTTPException(status_code=400, detail="El ticket de compra no contiene productos")

    now = datetime.now(timezone.utc)
    total_cost = 0.0

    purchase = models.Purchase(
        provider_id=provider.id,
        purchase_date=now,
        ticket_number=data.ticket_number or f"TICK-{int(now.timestamp())}",
        total_cost=0.0,
        notes=data.notes,
        user_id=data.user_id
    )
    db.add(purchase)
    db.commit()
    db.refresh(purchase)

    for item in data.items:
        prod = db.query(models.Product).filter(models.Product.id == item.product_id).first()
        if not prod:
            continue

        item_subtotal = round(item.quantity * item.unit_price_paid, 2)
        total_cost += item_subtotal

        p_item = models.PurchaseItem(
            purchase_id=purchase.id,
            product_id=prod.id,
            quantity=item.quantity,
            unit_price_paid=item.unit_price_paid,
            subtotal=item_subtotal
        )
        db.add(p_item)

        prod.current_stock = (prod.current_stock or 0.0) + item.quantity
        prod.unit_price = item.unit_price_paid
        prod.last_purchased_price = item.unit_price_paid
        prod.last_purchased_date = now

    purchase.total_cost = round(total_cost, 2)

    audit = models.AuditLog(
        user_id=data.user_id,
        action="REGISTER_PURCHASE",
        details=f"Compra registrada en {provider.name}. Ticket: {purchase.ticket_number}. Total: ${purchase.total_cost:.2f} ({len(data.items)} insumos)"
    )
    db.add(audit)
    db.commit()

    return {
        "message": f"¡Compra registrada en {provider.name} con éxito!",
        "purchase_id": purchase.id,
        "ticket_number": purchase.ticket_number,
        "total_cost": purchase.total_cost,
        "items_count": len(data.items)
    }


@app.get("/api/purchases")
def get_purchases(limit: int = 50, db: Session = Depends(get_db)):
    purchases = db.query(models.Purchase).order_by(models.Purchase.purchase_date.desc()).limit(limit).all()
    result = []

    for p in purchases:
        result.append({
            "id": p.id,
            "provider_name": p.provider.name if p.provider else "Desconocido",
            "purchase_date": p.purchase_date.strftime("%Y-%m-%d %H:%M"),
            "ticket_number": p.ticket_number or "S/N",
            "total_cost": p.total_cost or 0.0,
            "notes": p.notes or "",
            "items_count": len(p.items),
            "items": [
                {
                    "product_id": it.product_id,
                    "product_name": it.product.name if it.product else "N/A",
                    "category": it.product.category if it.product else "",
                    "quantity": it.quantity,
                    "unit": it.product.unit_measure if it.product else "",
                    "unit_price_paid": it.unit_price_paid,
                    "subtotal": it.subtotal
                }
                for it in p.items
            ]
        })
    return result


# --- CORTE DIARIO Y CONSUMO DE COCINA ---

@app.post("/api/consume-stock")
def consume_stock(data: StockConsume, db: Session = Depends(get_db)):
    product = db.query(models.Product).filter(models.Product.id == data.product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    
    current = product.current_stock or 0.0
    consumed = data.consumed_amount
    new_stock = max(0.0, current - consumed)
    
    unit_cost = product.unit_price or 0.0
    cost_total = round(consumed * unit_cost, 2)
    now = datetime.now(timezone.utc)
    reason = data.reason_type or "MERMA"

    product.current_stock = new_stock

    clog = models.ConsumptionLog(
        product_id=product.id,
        quantity=consumed,
        unit_price_at_moment=unit_cost,
        total_cost=cost_total,
        timestamp=now,
        reason_type=reason,
        shift_notes=data.shift_notes or f"Salida por {reason}",
        user_id=data.user_id
    )
    db.add(clog)

    audit = models.AuditLog(
        user_id=data.user_id,
        action="RECORD_WASTE_CONSUMPTION",
        details=f"Salida [{reason}]: {consumed} {product.unit_measure} de '{product.name}'. Costo: ${cost_total:.2f}. Motivo: {data.shift_notes or 'S/N'}. Stock restante: {new_stock}"
    )
    db.add(audit)
    db.commit()
    db.refresh(clog)

    return {
        "message": f"Salida registrada con éxito para {product.name}",
        "log_id": clog.id,
        "consumed_amount": consumed,
        "reason_type": reason,
        "current_stock": new_stock,
        "cost_total": cost_total
    }


@app.delete("/api/consumptions/{log_id}")
def delete_consumption(log_id: int, db: Session = Depends(get_db)):
    """Revierte una salida o merma registrada por error, regresando el stock al almacén."""
    log = db.query(models.ConsumptionLog).filter(models.ConsumptionLog.id == log_id).first()
    if not log:
        raise HTTPException(status_code=404, detail="Registro de consumo no encontrado")

    prod = log.product
    if prod:
        prod.current_stock = round((prod.current_stock or 0.0) + log.quantity, 3)

    audit = models.AuditLog(
        action="CANCEL_WASTE_CONSUMPTION",
        details=f"Cancelación de merma/salida id={log.id} ({log.reason_type}): {log.quantity} de '{prod.name if prod else 'N/A'}' revertido a stock."
    )
    db.add(audit)
    db.delete(log)
    db.commit()

    return {"message": "Registro revertido y stock reintegrado con éxito", "restored_stock": prod.current_stock if prod else 0}


@app.post("/api/update-stock")
def update_stock(data: StockUpdate, db: Session = Depends(get_db)):
    product = db.query(models.Product).filter(models.Product.id == data.product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Producto no encontrado")

    old_stock = product.current_stock or 0.0
    product.current_stock = max(0.0, data.new_stock)
    
    audit = models.AuditLog(
        action="UPDATE_STOCK_DIRECT",
        details=f"Ajuste manual de stock para {product.name}: {old_stock} -> {product.current_stock}"
    )
    db.add(audit)
    db.commit()
    return {"message": "Stock actualizado", "current_stock": product.current_stock}


@app.get("/api/consumptions")
def get_consumptions(limit: int = 50, db: Session = Depends(get_db)):
    logs = db.query(models.ConsumptionLog).order_by(models.ConsumptionLog.timestamp.desc()).limit(limit).all()
    return [
        {
            "id": log.id,
            "product_name": log.product.name if log.product else "N/A",
            "category": log.product.category if log.product else "",
            "quantity": log.quantity,
            "unit": log.product.unit_measure if log.product else "",
            "unit_price": log.unit_price_at_moment,
            "total_cost": log.total_cost,
            "timestamp": log.timestamp.strftime("%Y-%m-%d %H:%M"),
            "reason_type": getattr(log, 'reason_type', None) or "MERMA",
            "shift_notes": log.shift_notes
        }
        for log in logs
    ]


# --- INTEGRACIÓN CORTE POS, RECETAS Y ESCANDALLOS ---

def process_corte_parsed(header: Dict[str, Any], raw_products: List[Dict[str, Any]], db: Session):
    """Cruza los platillos del reporte de corte con el catálogo de recetas y calcula el descuento de almacén."""
    menu_items = {m.pos_name.strip().upper(): m for m in db.query(models.MenuItem).all()}
    
    parsed_sold_items = []
    ingredients_map = {}
    unmapped_items = []
    total_cogs = 0.0
    total_qty = 0.0

    for p in raw_products:
        p_name = p["name"].strip()
        p_name_upper = p_name.upper()
        qty = p["quantity"]
        total_p = p["total_amount"]
        total_qty += qty

        # Buscar en menú
        menu = menu_items.get(p_name_upper)
        dish_cogs = 0.0

        if menu and menu.recipes:
            for rec in menu.recipes:
                ing_prod = rec.product
                if not ing_prod:
                    continue
                needed_qty = rec.quantity_needed * qty
                cost = round(needed_qty * (ing_prod.unit_price or 0.0), 2)
                dish_cogs += cost

                if ing_prod.id not in ingredients_map:
                    ingredients_map[ing_prod.id] = {
                        "product_id": ing_prod.id,
                        "name": ing_prod.name,
                        "category": ing_prod.category,
                        "unit": ing_prod.unit_measure,
                        "current_stock": ing_prod.current_stock or 0.0,
                        "unit_price": ing_prod.unit_price or 0.0,
                        "quantity_to_deduct": 0.0,
                        "total_cost": 0.0
                    }
                ingredients_map[ing_prod.id]["quantity_to_deduct"] = round(ingredients_map[ing_prod.id]["quantity_to_deduct"] + needed_qty, 3)
                ingredients_map[ing_prod.id]["total_cost"] = round(ingredients_map[ing_prod.id]["total_cost"] + cost, 2)
        else:
            unmapped_items.append(p_name)

        total_cogs += dish_cogs
        parsed_sold_items.append({
            "name": p_name,
            "category": menu.category if menu else "GENERAL",
            "quantity": qty,
            "unit_price": p["unit_price"],
            "total_amount": total_p,
            "has_recipe": bool(menu and menu.recipes),
            "cogs_cost": round(dish_cogs, 2)
        })

    total_sales = header["ventas_totales"]
    gross_profit = round(total_sales - total_cogs, 2)
    food_cost_pct = round((total_cogs / total_sales * 100), 1) if total_sales > 0 else 0.0

    return {
        "summary": {
            "date_str": header["date_str"],
            "total_sales": total_sales,
            "cash_sales": header["efectivo"],
            "card_sales": header["tarjeta"],
            "transfer_sales": header["transferencias"],
            "tips_cash": header["propinas_efectivo"],
            "tips_card": header["propinas_tarjeta"],
            "tips_total": header["propinas_totales"],
            "cash_balance": header["saldo_caja"],
            "total_items_sold": int(total_qty),
            "total_cogs": round(total_cogs, 2),
            "gross_profit": gross_profit,
            "food_cost_percentage": food_cost_pct,
            "unmapped_count": len(unmapped_items)
        },
        "sold_items": parsed_sold_items,
        "ingredients_to_deduct": list(ingredients_map.values()),
        "unmapped_items": list(set(unmapped_items))
    }


@app.post("/api/sales-cut/parse-text")
def parse_sales_cut_text(payload: ParseTextRequest, db: Session = Depends(get_db)):
    """Parsea el texto crudo pegado desde el Excel del POS."""
    f = io.StringIO(payload.text.strip())
    reader = csv.reader(f)
    rows = list(reader)
    header, products = parse_corte_rows(rows)
    return process_corte_parsed(header, products, db)


@app.post("/api/sales-cut/upload-file")
async def parse_sales_cut_file(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Parsea un archivo subido (.xlsx, .xls o .csv) leyendo todas las hojas disponibles."""
    contents = await file.read()
    filename = file.filename.lower()
    all_rows = []
    parsed_excel = False

    # 1. Intentar leer como Excel con openpyxl (soporta .xlsx y .xls generados por POS que son XLSX por dentro)
    is_excel = filename.endswith(".xlsx") or filename.endswith(".xls") or contents.startswith(b"PK\x03\x04")
    if is_excel or contents.startswith(b"PK\x03\x04"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                for r in ws.iter_rows(values_only=True):
                    if any(r):
                        all_rows.append([str(c) if c is not None else "" for c in r])
            parsed_excel = True
        except Exception:
            # Fallback a xlrd si es un archivo .xls binario clásico (BIFF8)
            try:
                import xlrd
                wb_xls = xlrd.open_workbook(file_contents=contents)
                for s_name in wb_xls.sheet_names():
                    ws = wb_xls.sheet_by_name(s_name)
                    for row_idx in range(ws.nrows):
                        r_vals = ws.row_values(row_idx)
                        if any(r_vals):
                            all_rows.append([str(c) if c is not None else "" for c in r_vals])
                parsed_excel = True
            except Exception:
                parsed_excel = False

    # 2. Si no es un Excel válido o falló, intentar parsearlo como CSV / texto con varias codificaciones
    if not parsed_excel:
        for enc in ["utf-8-sig", "latin-1", "cp1252", "utf-8"]:
            try:
                text_str = contents.decode(enc)
                f = io.StringIO(text_str)
                try:
                    dialect = csv.Sniffer().sniff(text_str[:2048])
                    reader = csv.reader(f, dialect)
                except Exception:
                    f.seek(0)
                    reader = csv.reader(f)
                all_rows = list(reader)
                if all_rows:
                    break
            except Exception:
                continue

    header, products = parse_corte_rows(all_rows)
    return process_corte_parsed(header, products, db)


@app.post("/api/sales-cut/confirm")
def confirm_sales_cut(data: ConfirmCutRequest, db: Session = Depends(get_db)):
    """Aplica el corte de caja: descuenta almacén, genera bitácoras y guarda histórico financiero."""
    now = datetime.now(timezone.utc)

    # 1. Crear registro SalesCut
    sales_cut = models.SalesCut(
        cut_date_str=data.cut_date_str,
        cut_date=now,
        total_sales=data.total_sales,
        cash_sales=data.cash_sales,
        card_sales=data.card_sales,
        transfer_sales=data.transfer_sales,
        tips_cash=data.tips_cash,
        tips_card=data.tips_card,
        tips_total=data.tips_total,
        cash_balance=data.cash_balance,
        total_items_sold=data.total_items_sold,
        total_cogs=data.total_cogs,
        gross_profit=data.gross_profit,
        food_cost_percentage=data.food_cost_percentage
    )
    db.add(sales_cut)
    db.commit()
    db.refresh(sales_cut)

    # 2. Guardar partidas de platillos vendidos
    for item in data.sold_items:
        s_item = models.SalesCutItem(
            sales_cut_id=sales_cut.id,
            product_name=item.name,
            category=item.category if hasattr(item, 'category') else "GENERAL",
            quantity=item.quantity,
            unit_price=item.unit_price,
            total_amount=item.total_amount,
            cogs_cost=item.cogs_cost
        )
        db.add(s_item)

    # 3. Descontar insumos del almacén y registrar bitácora de consumo
    for ing in data.ingredients_to_deduct:
        prod = db.query(models.Product).filter(models.Product.id == ing.product_id).first()
        if prod:
            current_s = prod.current_stock or 0.0
            new_s = max(0.0, round(current_s - ing.quantity_to_deduct, 3))
            prod.current_stock = new_s

            clog = models.ConsumptionLog(
                product_id=prod.id,
                quantity=ing.quantity_to_deduct,
                unit_price_at_moment=prod.unit_price or 0.0,
                total_cost=ing.cost,
                timestamp=now,
                reason_type="VENTA_POS",
                shift_notes=f"Corte POS: {data.cut_date_str}"
            )
            db.add(clog)

    audit = models.AuditLog(
        action="APPLY_POS_CUT",
        details=f"Corte POS aplicado: Ventas ${data.total_sales:.2f}, Costo Insumos ${data.total_cogs:.2f} ({data.food_cost_percentage}% Food Cost). Insumos descontados: {len(data.ingredients_to_deduct)}"
    )
    db.add(audit)
    db.commit()

    return {
        "message": "¡Corte de ventas procesado y almacén descontado con éxito!",
        "sales_cut_id": sales_cut.id,
        "total_sales": sales_cut.total_sales,
        "total_cogs": sales_cut.total_cogs,
        "gross_profit": sales_cut.gross_profit,
        "food_cost_percentage": sales_cut.food_cost_percentage
    }


@app.get("/api/sales-cuts")
def get_sales_cuts(limit: int = 30, db: Session = Depends(get_db)):
    """Historial de cortes de ventas procesados."""
    cuts = db.query(models.SalesCut).order_by(models.SalesCut.cut_date.desc()).limit(limit).all()
    return [
        {
            "id": c.id,
            "cut_date_str": c.cut_date_str,
            "cut_date": c.cut_date.strftime("%Y-%m-%d %H:%M"),
            "total_sales": c.total_sales,
            "cash_sales": c.cash_sales,
            "card_sales": c.card_sales,
            "transfer_sales": c.transfer_sales,
            "tips_total": c.tips_total,
            "cash_balance": c.cash_balance,
            "total_items_sold": c.total_items_sold,
            "total_cogs": c.total_cogs,
            "gross_profit": c.gross_profit,
            "food_cost_percentage": c.food_cost_percentage,
            "items_count": len(c.items)
        }
        for c in cuts
    ]


# --- GESTIÓN DE RECETAS Y ESCANDALLOS ---

@app.get("/api/menu-items")
def get_menu_items(db: Session = Depends(get_db)):
    items = db.query(models.MenuItem).order_by(models.MenuItem.pos_name.asc()).all()
    result = []
    for m in items:
        # Calcular costo de receta
        cost = sum(r.quantity_needed * (r.product.unit_price or 0.0) for r in m.recipes if r.product)
        margin = m.sale_price - cost if m.sale_price else 0.0
        fc_pct = (cost / m.sale_price * 100) if m.sale_price > 0 else 0.0

        result.append({
            "id": m.id,
            "pos_name": m.pos_name,
            "category": m.category,
            "sale_price": m.sale_price,
            "recipe_cost": round(cost, 2),
            "margin": round(margin, 2),
            "food_cost_percentage": round(fc_pct, 1),
            "recipes_count": len(m.recipes),
            "recipes": [
                {
                    "product_id": r.product_id,
                    "product_name": r.product.name if r.product else "N/A",
                    "unit": r.product.unit_measure if r.product else "",
                    "unit_price": r.product.unit_price if r.product else 0.0,
                    "quantity_needed": r.quantity_needed,
                    "cost": round(r.quantity_needed * (r.product.unit_price or 0.0), 2)
                }
                for r in m.recipes
            ]
        })
    return result


@app.post("/api/menu-items")
def save_menu_item(data: MenuItemCreate, db: Session = Depends(get_db)):
    menu = db.query(models.MenuItem).filter(models.MenuItem.pos_name == data.pos_name).first()
    if not menu:
        menu = models.MenuItem(
            pos_name=data.pos_name,
            category=data.category,
            sale_price=data.sale_price
        )
        db.add(menu)
        db.commit()
        db.refresh(menu)
    else:
        menu.category = data.category
        menu.sale_price = data.sale_price

    # Reemplazar recetas
    db.query(models.RecipeItem).filter(models.RecipeItem.menu_item_id == menu.id).delete()
    for r in data.recipes:
        rec = models.RecipeItem(
            menu_item_id=menu.id,
            product_id=r.product_id,
            quantity_needed=r.quantity_needed
        )
        db.add(rec)
    db.commit()
    return {"message": f"Receta de '{menu.pos_name}' guardada con éxito", "id": menu.id}


# --- CENTRO DE MANDO & DASHBOARD DE MÉTRICAS EJECUTIVAS 360° (AMPLIADO) ---

@app.get("/api/dashboard-metrics")
def get_dashboard_metrics(days: int = 30, db: Session = Depends(get_db)):
    all_products = db.query(models.Product).all()
    now = datetime.now(timezone.utc)
    start_date = now - timedelta(days=days)

    # 1. KPIs Generales de Almacén
    total_inventory_value = 0.0
    estimated_restock_budget = 0.0
    critical_count = 0
    low_count = 0
    optimal_count = 0

    category_data = {cat: {"inventory_value": 0.0, "restock_budget": 0.0, "items_count": 0, "critical": 0} for cat in CATEGORIES_LIST}
    provider_data = {}

    for prod in all_products:
        stock = prod.current_stock or 0.0
        ideal = prod.ideal_stock or 0.0
        price = prod.unit_price or 0.0
        cat = prod.category or "Abarrotes y Harinas"
        prov = prod.provider.name if prod.provider else "Otros"

        val = stock * price
        total_inventory_value += val

        deficit = max(0.0, ideal - stock)
        deficit_cost = deficit * price
        estimated_restock_budget += deficit_cost

        if stock <= 0:
            critical_count += 1
            if cat in category_data:
                category_data[cat]["critical"] += 1
        elif stock < ideal:
            low_count += 1
        else:
            optimal_count += 1

        if cat not in category_data:
            category_data[cat] = {"inventory_value": 0.0, "restock_budget": 0.0, "items_count": 0, "critical": 0}
        category_data[cat]["inventory_value"] += val
        category_data[cat]["restock_budget"] += deficit_cost
        category_data[cat]["items_count"] += 1

        if prov not in provider_data:
            provider_data[prov] = {
                "name": prov,
                "inventory_value": 0.0,
                "restock_budget": 0.0,
                "deficit_items": 0,
                "total_items": 0
            }
        provider_data[prov]["inventory_value"] += val
        provider_data[prov]["restock_budget"] += deficit_cost
        provider_data[prov]["total_items"] += 1
        if stock < ideal:
            provider_data[prov]["deficit_items"] += 1

    # 2. Compras Históricas Reales
    purchases_in_period = db.query(models.Purchase).filter(models.Purchase.purchase_date >= start_date).all()
    total_purchases_month = sum((p.total_cost or 0.0) for p in purchases_in_period)

    category_purchases = {cat: 0.0 for cat in CATEGORIES_LIST}
    provider_purchases = {}
    for p in purchases_in_period:
        prov_name = p.provider.name if p.provider else "Otros"
        provider_purchases[prov_name] = provider_purchases.get(prov_name, 0.0) + (p.total_cost or 0.0)
        for item in p.items:
            cat = item.product.category if item.product else "Abarrotes y Harinas"
            category_purchases[cat] = category_purchases.get(cat, 0.0) + (item.subtotal or 0.0)

    # 3. Consumo de Cocina en el Periodo
    consumptions_in_period = db.query(models.ConsumptionLog).filter(models.ConsumptionLog.timestamp >= start_date).all()
    total_consumption_month = sum((c.total_cost or 0.0) for c in consumptions_in_period)

    category_consumption = {cat: 0.0 for cat in CATEGORIES_LIST}
    product_consumption_map = {}
    for c in consumptions_in_period:
        p_name = c.product.name if c.product else "N/A"
        cat = c.product.category if c.product else "Abarrotes y Harinas"
        category_consumption[cat] = category_consumption.get(cat, 0.0) + (c.total_cost or 0.0)
        
        if p_name not in product_consumption_map:
            product_consumption_map[p_name] = {
                "name": p_name,
                "category": cat,
                "total_quantity": 0.0,
                "unit": c.product.unit_measure if c.product else "",
                "total_cost": 0.0
            }
        product_consumption_map[p_name]["total_quantity"] += c.quantity
        product_consumption_map[p_name]["total_cost"] += c.total_cost

    top_consumed = sorted(product_consumption_map.values(), key=lambda x: x["total_cost"], reverse=True)[:5]

    # 4. Métricas de Ventas POS (Cortes Diarios)
    sales_cuts_in_period = db.query(models.SalesCut).filter(models.SalesCut.cut_date >= start_date).all()
    total_sales_period = sum((c.total_sales or 0.0) for c in sales_cuts_in_period)
    total_cogs_period = sum((c.total_cogs or 0.0) for c in sales_cuts_in_period)
    gross_profit_period = sum((c.gross_profit or 0.0) for c in sales_cuts_in_period)
    cash_sales = sum((c.cash_sales or 0.0) for c in sales_cuts_in_period)
    card_sales = sum((c.card_sales or 0.0) for c in sales_cuts_in_period)
    transfer_sales = sum((c.transfer_sales or 0.0) for c in sales_cuts_in_period)
    tips_total = sum((c.tips_total or 0.0) for c in sales_cuts_in_period)
    avg_food_cost = round((total_cogs_period / total_sales_period * 100), 1) if total_sales_period > 0 else 0.0

    # Top Platillos Más Vendidos en Cortes POS
    top_dishes_query = db.query(
        models.SalesCutItem.product_name,
        func.sum(models.SalesCutItem.quantity).label("total_qty"),
        func.sum(models.SalesCutItem.total_amount).label("total_money")
    ).group_by(models.SalesCutItem.product_name).order_by(func.sum(models.SalesCutItem.total_amount).desc()).limit(8).all()

    top_dishes = [
        {"name": row[0], "quantity": row[1], "total_sales": round(row[2], 2)}
        for row in top_dishes_query
    ]

    # 5. Tendencia Temporal (Compras vs Consumo por Fecha)
    date_map = {}
    for i in range(days, -1, -1):
        d_str = (now - timedelta(days=i)).strftime("%Y-%m-%d")
        date_map[d_str] = {"date": d_str, "compras": 0.0, "consumo": 0.0, "ventas": 0.0}

    for p in purchases_in_period:
        d_str = p.purchase_date.strftime("%Y-%m-%d")
        if d_str in date_map:
            date_map[d_str]["compras"] += p.total_cost or 0.0

    for c in consumptions_in_period:
        d_str = c.timestamp.strftime("%Y-%m-%d")
        if d_str in date_map:
            date_map[d_str]["consumo"] += c.total_cost or 0.0

    for sc in sales_cuts_in_period:
        d_str = sc.cut_date.strftime("%Y-%m-%d")
        if d_str in date_map:
            date_map[d_str]["ventas"] += sc.total_sales or 0.0

    timeline = list(date_map.values())

    category_chart = [
        {
            "category": cat,
            "inventory_value": round(category_data[cat]["inventory_value"], 2),
            "restock_budget": round(category_data[cat]["restock_budget"], 2),
            "historical_spent": round(category_purchases.get(cat, 0.0), 2),
            "consumption_spent": round(category_consumption.get(cat, 0.0), 2),
            "critical_count": category_data[cat]["critical"]
        }
        for cat in CATEGORIES_LIST
    ]

    provider_chart = [
        {
            "provider": p_info["name"],
            "inventory_value": round(p_info["inventory_value"], 2),
            "restock_budget": round(p_info["restock_budget"], 2),
            "historical_spent": round(provider_purchases.get(p_info["name"], 0.0), 2),
            "deficit_items": p_info["deficit_items"],
            "total_items": p_info["total_items"]
        }
        for p_info in provider_data.values()
    ]

    critical_items = [
        {
            "product_id": p.id,
            "name": p.name,
            "category": p.category,
            "provider": p.provider.name if p.provider else "N/A",
            "current_stock": p.current_stock or 0.0,
            "ideal_stock": p.ideal_stock or 0.0,
            "deficit": max(0.0, (p.ideal_stock or 0.0) - (p.current_stock or 0.0)),
            "unit": p.unit_measure,
            "estimated_cost": round(max(0.0, (p.ideal_stock or 0.0) - (p.current_stock or 0.0)) * (p.unit_price or 0.0), 2)
        }
        for p in all_products if (p.current_stock or 0.0) <= 0 or (p.current_stock or 0.0) < (p.ideal_stock or 0.0) * 0.3
    ]
    critical_items = sorted(critical_items, key=lambda x: (x["current_stock"] <= 0, x["deficit"]), reverse=True)[:10]

    waste_cost_period = sum((c.total_cost or 0.0) for c in consumptions_in_period if getattr(c, 'reason_type', None) == 'MERMA')
    cleaning_cost_period = sum((c.total_cost or 0.0) for c in consumptions_in_period if getattr(c, 'reason_type', None) == 'LIMPIEZA')
    staff_food_cost_period = sum((c.total_cost or 0.0) for c in consumptions_in_period if getattr(c, 'reason_type', None) == 'COMIDA_PERSONAL')

    return {
        "kpis": {
            "total_inventory_value": round(total_inventory_value, 2),
            "estimated_restock_budget": round(estimated_restock_budget, 2),
            "total_purchases_period": round(total_purchases_month, 2),
            "total_consumption_period": round(total_consumption_month, 2),
            "waste_cost_period": round(waste_cost_period, 2),
            "cleaning_cost_period": round(cleaning_cost_period, 2),
            "staff_food_cost_period": round(staff_food_cost_period, 2),
            "critical_items_count": critical_count,
            "low_stock_items_count": low_count,
            "optimal_items_count": optimal_count,
            "total_items_count": len(all_products),
            # Nuevos KPIs financieros de ventas POS
            "total_sales_period": round(total_sales_period, 2),
            "total_cogs_period": round(total_cogs_period, 2),
            "gross_profit_period": round(gross_profit_period, 2),
            "avg_food_cost_percentage": avg_food_cost,
            "cuts_count": len(sales_cuts_in_period)
        },
        "payment_methods": {
            "cash": round(cash_sales, 2),
            "card": round(card_sales, 2),
            "transfer": round(transfer_sales, 2),
            "tips": round(tips_total, 2)
        },
        "top_dishes": top_dishes,
        "category_chart": category_chart,
        "provider_chart": provider_chart,
        "timeline": timeline,
        "top_consumed": top_consumed,
        "critical_items": critical_items
    }


# --- PROVEEDORES Y PRODUCTOS (CRUD) ---

@app.get("/api/providers")
def get_providers(db: Session = Depends(get_db)):
    return db.query(models.Provider).order_by(models.Provider.name.asc()).all()

@app.post("/api/providers")
def create_provider(data: ProviderCreate, db: Session = Depends(get_db)):
    exists = db.query(models.Provider).filter(models.Provider.name == data.name).first()
    if exists:
        raise HTTPException(status_code=400, detail="Ese proveedor ya existe")
    new_prov = models.Provider(name=data.name)
    db.add(new_prov)
    db.commit()
    db.refresh(new_prov)
    return new_prov

@app.post("/api/products")
def create_product(data: ProductCreate, db: Session = Depends(get_db)):
    new_prod = models.Product(
        name=data.name,
        category=data.category,
        provider_id=data.provider_id,
        ideal_stock=data.ideal_stock,
        current_stock=data.current_stock,
        unit_measure=data.unit_measure,
        unit_price=data.unit_price,
        last_purchased_price=data.unit_price,
        last_purchased_date=datetime.now(timezone.utc)
    )
    db.add(new_prod)
    db.commit()
    db.refresh(new_prod)
    return new_prod

@app.put("/api/products/{product_id}")
def update_product(product_id: int, data: ProductUpdate, db: Session = Depends(get_db)):
    product = db.query(models.Product).filter(models.Product.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    
    product.name = data.name
    product.category = data.category
    product.ideal_stock = data.ideal_stock
    product.unit_measure = data.unit_measure
    product.unit_price = data.unit_price
    if data.provider_id:
        product.provider_id = data.provider_id
    db.commit()
    return {"message": "Producto actualizado con éxito"}

@app.delete("/api/products/{product_id}")
def delete_product(product_id: int, db: Session = Depends(get_db)):
    product = db.query(models.Product).filter(models.Product.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    
    db.delete(product)
    db.commit()
    return {"message": "Producto eliminado con éxito"}